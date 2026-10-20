from pathlib import Path
from openpyxl import load_workbook
from pypdf import PdfReader
import csv
root=Path('tmp/stage11-exports')
for width in (320,390,768,1280):
    workbook=load_workbook(root/f'cashbook-{width}.xlsx',data_only=False)
    sheet=workbook['Expense cashbook']
    assert sheet.max_row==1008
    details=dict(workbook['Export details'].values)
    assert details['Campus']=='senior' and details['Source records']=='1007'
    assert details['pending USD']=='1258.75'
    assert sum(row[6].value for row in list(sheet.rows)[1:])==1258.75
    assert sheet['I2'].value=='0000123' and sheet['I2'].data_type=='s'
    assert sheet['E2'].value==' =SUM(1,2)' and sheet['E2'].data_type=='s'
    assert all(cell.data_type!='f' for row in sheet for cell in row)
    with (root/f'cashbook-{width}.csv').open(encoding='utf-8-sig',newline='') as stream:
        rows=list(csv.reader(stream))
    data=[row for row in rows if row and row[0].startswith('expense-')]
    assert len(data)==1007
    assert sum(float(row[6]) for row in data)==1258.75
    assert data[0][4]=="' =SUM(1,2)" and data[0][8]=='0000123'
    print(f'PASS {width}px actual CSV/XLSX: every row, numeric totals, literal identifiers and formula safety')
stock=load_workbook(root/'stock.xlsx',data_only=False)['Stock and sizes']
assert stock.max_row==3
assert [stock.cell(row,10).value for row in (2,3)]==[2,3]
assert stock['H2'].value=='0007' and stock['H2'].data_type=='s'
sales=load_workbook(root/'sales.xlsx',data_only=False)
assert sales['POS receipts']['G2'].value==15 and sales['POS receipts']['H2'].value==20
assert sales['Saved order lines']['C2'].value=='School uniform shirt with a longer description'
assert 'Renamed current shirt' not in sales['Saved order lines']['C2'].value
orders=load_workbook(root/'orders.xlsx',data_only=False)
assert orders['Saved order lines']['C2'].value=='Original tie'
assert orders['Saved order lines']['G2'].value==10
print('PASS real stock sizes, literal SKU, tender versus sales total, and historical saved order lines')
for name, count in [('filtered',3),('selected',1)]:
    book=load_workbook(root/f'fees-{name}.xlsx',data_only=False)
    sheet=book['Fee balances']
    assert sheet.max_row==count+1
    assert sum(sheet.cell(row,7).value for row in range(2,count+2))==135*count
    assert dict(book['Export details'].values)['Learners']==count
    assert all(sheet.cell(row,2).data_type=='s' for row in range(2,count+2))
receipt=PdfReader(root/'receipt.pdf')
assert len(receipt.pages)==1
text=receipt.pages[0].extract_text()
assert text.count('TEST-RECEIPT-001')==1
assert 'Received $20.00' in text and 'Change $5.00' in text
assert 'Download records' not in text
print('PASS filtered/selected fee totals and one-page receipt PDF without duplicate receipts or portal controls')
