"""Verify real browser downloads, not mocked XLSX/PDF generation."""
from pathlib import Path
import csv
import re
from openpyxl import load_workbook
from pypdf import PdfReader

root = Path('tmp/stage10-exports')
for width in (320, 390, 667, 1280):
    count = 1007 if width == 390 else 2
    with (root / f'register-{width}.csv').open(encoding='utf-8-sig', newline='') as stream:
        rows = list(csv.reader(stream))
    assert len(rows) == count + 1, (width, 'CSV row count', len(rows))
    workbook = load_workbook(root / f'register-{width}.xlsx', data_only=False)
    assert workbook.sheetnames == ['Learners', 'Export details']
    sheet = workbook['Learners']
    assert sheet.max_row == count + 1 and sheet.max_column == 13
    assert all(cell.data_type != 'f' for row in sheet for cell in row), 'Learner data must never become formulas'
    details = dict(workbook['Export details'].values)
    assert details['Campus'] == 'senior' and details['Learners'] == str(count)
    if width == 390:
        assert sheet['A2'].value == '00000000' and sheet['A2'].data_type == 's'
        assert sheet['B2'].value == ' \t=HYPERLINK("bad")'
        assert sheet['K2'].value == '+263770000001' and sheet['K2'].data_type == 's'
        assert sheet['F2'].value == '001-002'
        assert sheet['L2'].value == 'House "A"\nHarare'
        assert sheet['D2'].value == '13'
        assert rows[1][1].startswith("' \t=") and rows[1][10].startswith("'+263")
        assert rows[-1][0] == '00001006'
    pdf = PdfReader(root / f'register-{width}.pdf')
    for page in pdf.pages:
        assert float(page.mediabox.width) > float(page.mediabox.height), 'Register must be A4 landscape'
        assert abs(float(page.mediabox.width) - 841.89) < 2
        assert 'Admission no.' in re.sub(r'\s+', ' ', page.extract_text()), 'Repeat register column headers on every page'
    text = '\n'.join(page.extract_text() for page in pdf.pages)
    assert 'reliance learning centre' in text.lower() and f'{count} learners' in text
    assert 'Download CSV' not in text and 'Preview learner list' not in text and 'Log out' not in text
    if width == 390:
        assert len(pdf.pages) > 1
        assert '00000000' in text and '00001006' in text
        assert all(str(n).zfill(8) in text for n in range(1007)), 'Every admission identifier must reach the PDF'
    print(f'PASS {width}px: {count} CSV/XLSX rows, literal text cells, {len(pdf.pages)} landscape PDF pages, repeated headers and complete content')
