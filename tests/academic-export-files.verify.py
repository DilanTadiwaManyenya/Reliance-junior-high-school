from pathlib import Path
import csv,re
from openpyxl import load_workbook
from pypdf import PdfReader
root=Path('tmp/stage12-exports')
for width in (320,390,768,1280):
    book=load_workbook(root/f'batch-{width}.xlsx',data_only=False)
    assert book.sheetnames==['Report contexts','Subjects','Components','Development','Report comments','Export details']
    sheet=book['Subjects']
    assert sheet.max_row==1008
    assert sheet['I1008'].value=='Subject 1006'
    assert sheet['J2'].value==0 and sheet['L2'].value==0
    assert sheet['M2'].value==' =SUM(1,2)' and sheet['M2'].data_type=='s'
    assert all(cell.data_type!='f' for row in sheet for cell in row)
    with(root/f'batch-{width}.csv').open(encoding='utf-8-sig',newline='') as stream: rows=list(csv.reader(stream))
    marks=[row for row in rows if len(row)==15 and row[8].startswith('Subject ')]
    assert len(marks)==1007 and marks[0][12]=="' =SUM(1,2)"
    pdf=PdfReader(root/f'batch-{width}.pdf');text='\n'.join(page.extract_text() for page in pdf.pages)
    assert len(pdf.pages)>1
    assert len(set(re.findall(r'Subject\s+(\d{4})',text)))==1007
    assert sum('Admission ID' in page.extract_text() for page in pdf.pages)>1
    assert all(float(page.mediabox.width)>float(page.mediabox.height) for page in pdf.pages)
    assert 'Prepare academic export' not in text
    print(f'PASS {width}px: all 1007 marks in real CSV/XLSX and {len(pdf.pages)} landscape PDF pages, repeated headers, safe literal text and zero marks')
parent=load_workbook(root/'parent-historical.xlsx',data_only=False)
assert parent['Report contexts']['C2'].value=='junior'
assert parent['Report contexts']['D2'].value=='Grade 1'
assert parent['Report contexts']['H2'].value=='legacy_inferred'
attendance=load_workbook(root/'attendance.xlsx',data_only=False)
assert attendance['Attendance'].max_row==2
details=dict(attendance['Export details'].values)
date=attendance['Attendance']['C2'].value
assert details['Start date']<=date<=details['End date']
assert 'Unmarked days are not absences' in details['Scope']
pdf=PdfReader(root/'attendance.pdf')
assert len(pdf.pages)==1 and 'Attendance' in pdf.pages[0].extract_text()
print('PASS parent historical class/provenance and attendance date range, single-page PDF and unmarked-day semantics')
